import csv
import io
import re
import unicodedata
from collections import defaultdict
from pathlib import Path

import streamlit as st
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


st.set_page_config(
    page_title="Wordwall Results Manager 2.0",
    page_icon="📚",
    layout="centered"
)

MAX_FILES = 60
MAX_FILE_MB = 8
MAX_TOTAL_MB = 120

STUDENT_PATTERN = re.compile(
    r"^\s*(\d+[A-Za-z]?)\s+(.+?)\s*$"
)

CORRECT = {
    "✔", "✓", "true", "correct",
    "correcto", "1", "yes", "si"
}

INCORRECT = {
    "✖", "✘", "✗", "✕", "❌",
    "false", "incorrect", "incorrecto", "0", "no"
}


def normalize(value):
    value = unicodedata.normalize("NFKD", str(value))
    characters = "".join(
        c for c in value
        if not unicodedata.combining(c)
    )
    return " ".join(characters.casefold().split())


def as_text(value):
    return "" if value is None else str(value).strip()


def read_rows(upload):
    raw = upload.getvalue()

    if upload.name.lower().endswith(".csv"):
        decoded = None

        for encoding in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                decoded = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                continue

        if decoded is None:
            raise ValueError("No se pudo decodificar el CSV")

        return csv.reader(
            io.StringIO(decoded, newline="")
        )

    book = load_workbook(
        io.BytesIO(raw),
        read_only=True,
        data_only=True
    )

    if "Results by student" not in book.sheetnames:
        book.close()
        raise ValueError(
            "Falta la hoja 'Results by student'"
        )

    sheet = book["Results by student"]

    def rows():
        try:
            for row in sheet.iter_rows(
                min_col=1,
                max_col=3,
                values_only=True
            ):
                yield row
        finally:
            book.close()

    return rows()

def parse_activity(upload):
    iterator = iter(read_rows(upload))

    try:
        first = next(iterator)
    except StopIteration:
        raise ValueError("El archivo está vacío")

    activity = as_text(first[0]) if first else ""

    if not activity:
        activity = Path(upload.name).stem

    scores = defaultdict(list)
    names = {}

    current = None
    correct = 0
    questions = 0

    def save_attempt():
        if (
            current is not None
            and questions > 0
            and len(scores[current]) < 3
        ):
            scores[current].append(
                100 * correct / questions
            )

    for row in iterator:
        a = as_text(row[0]) if len(row) > 0 else ""
        b = as_text(row[1]) if len(row) > 1 else ""
        c = as_text(row[2]) if len(row) > 2 else ""

        match = (
            STUDENT_PATTERN.fullmatch(a)
            if not b and not c
            else None
        )

        if match:
            save_attempt()

            group = match.group(1)
            name = match.group(2).strip()

            current = (group, normalize(name))
            names.setdefault(current, name)

            correct = 0
            questions = 0
            continue

        if current is None:
            continue

        mark = normalize(c)

        if mark in CORRECT or mark in INCORRECT:
            questions += 1

            if mark in CORRECT:
                correct += 1

    save_attempt()

    grades = {
        key: sum(values) / len(values)
        for key, values in scores.items()
        if values
    }

def make_report(activities, all_names, title):
    wb = Workbook()
    wb.remove(wb.active)

    by_group = defaultdict(list)

    for key, name in all_names.items():
        by_group[key[0]].append((key, name))

    header_fill = PatternFill(
        "solid",
        fgColor="174A69"
    )

    alternate_fill = PatternFill(
        "solid",
        fgColor="EAF1F5"
    )

    average_fill = PatternFill(
        "solid",
        fgColor="CFDDE6"
    )

    white_fill = PatternFill(
        "solid",
        fgColor="FFFFFF"
    )

    border = Border(
        left=Side(style="thin", color="D1D9DF"),
        right=Side(style="thin", color="D1D9DF"),
        top=Side(style="thin", color="D1D9DF"),
        bottom=Side(style="thin", color="D1D9DF")
    )

    groups = sorted(
        by_group,
        key=lambda g: (
            not g.isdigit(),
            g.zfill(10)
        )
    )

    for group in groups:
        sheet = wb.create_sheet(
            ("Grupo " + group)[:31]
        )

        headers = (
            ["Alumno"]
            + [name for name, _ in activities]
            + ["Promedio mensual"]
        )

        sheet.append(headers)
        sheet.freeze_panes = "B2"
        sheet.sheet_view.showGridLines = False
        sheet.row_dimensions[1].height = 35

        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = Font(
                bold=True,
                color="FFFFFF"
            )
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True
            )
            cell.border = border

        students = sorted(
            by_group[group],
            key=lambda item: normalize(item[1])
        )

        for row_num, (key, name) in enumerate(
            students,
            start=2
        ):
            values = [
                grades.get(key)
                for _, grades in activities
            ]

            average = sum(
                value if value is not None else 0
                for value in values
            ) / len(activities)

            sheet.append(
                [name]
                + [
                    round(value, 2)
                    if value is not None
                    else "-"
                    for value in values
                ]
                + [round(average, 2)]
            )

            striped = row_num % 2 == 1

            for col_num in range(
                1,
                len(headers) + 1
            ):
                cell = sheet.cell(
                    row_num,
                    col_num
                )

                is_average = (
                    col_num == len(headers)
                )

                if striped:
                    cell.fill = (
                        average_fill
                        if is_average
                        else alternate_fill
                    )
                else:
                    cell.fill = white_fill

                cell.border = border

                cell.font = Font(
                    bold=is_average,
                    color="263747"
                )

                cell.alignment = Alignment(
                    horizontal=(
                        "left"
                        if col_num == 1
                        else "center"
                    ),
                    vertical="center"
                )

                if (
                    col_num > 1
                    and isinstance(
                        cell.value,
                        (float, int)
                    )
                ):
                    cell.number_format = "0.00"

        for col_num, header in enumerate(
            headers,
            start=1
        ):
            if col_num == 1:
                width = min(
                    45,
                    max(
                        22,
                        max(
                            (
                                len(name)
                                for _, name in students
                            ),
                            default=0
                        ) + 3
                    )
                )
            else:
                width = min(
                    48,
                    max(17, len(header) + 3)
                )

            sheet.column_dimensions[
                get_column_letter(col_num)
            ].width = width

        sheet.auto_filter.ref = sheet.dimensions

    output = io.BytesIO()
    wb.save(output)

    return output.getvalue(), len(by_group)


def clean_filename(title):
    cleaned = normalize(title)

    cleaned = re.sub(
        r"[^a-z0-9_-]+",
        "_",
        cleaned
    ).strip("_")

    return (
        cleaned[:80] or "reporte_wordwall"
    ) + ".xlsx"

st.title("📚 Wordwall Results Manager")

st.info(
    "Organiza los resultados de tus alumnos de forma rápida y sencilla. "
    "Ingresa el título que deseas que tenga tu reporte, sube tus archivos de Excel "
    "y descarga el concentrado de calificaciones. "
)

st.caption(
    "Los archivos se procesan durante tu sesión; "
    "no se guardan en una base de datos de esta aplicación."
)

if "report_bytes" not in st.session_state:
    st.session_state.report_bytes = None
    st.session_state.report_name = None
    st.session_state.report_info = None

with st.form("upload_form", clear_on_submit=False):
    title = st.text_input(
        "Título del reporte",
        value="Reporte de Wordwall"
    )

    uploads = st.file_uploader(
        "Selecciona los archivos Excel o CSV",
        type=["xlsx", "csv"],
        accept_multiple_files=True
    )

    submitted = st.form_submit_button(
        "Generar reporte",
        type="primary",
        use_container_width=True
    )

if submitted:
    st.session_state.report_bytes = None
    st.session_state.report_name = None
    st.session_state.report_info = None

    if not uploads:
        st.warning(
            "Selecciona al menos un archivo."
        )

    elif len(uploads) > MAX_FILES:
        st.error(
            f"Puedes subir como máximo "
            f"{MAX_FILES} archivos por reporte."
        )

    elif any(
        u.size > MAX_FILE_MB * 1024 * 1024
        for u in uploads
    ):
        st.error(
            f"Cada archivo debe pesar como máximo "
            f"{MAX_FILE_MB} MB."
        )

    elif sum(
        u.size for u in uploads
    ) > MAX_TOTAL_MB * 1024 * 1024:
        st.error(
            f"El tamaño combinado no debe superar "
            f"{MAX_TOTAL_MB} MB."
        )

    else:
        activities = []
        names = {}
        seen = set()
        errors = []

        progress = st.progress(
            0,
            text="Procesando archivos..."
        )

        for i, upload in enumerate(
            uploads,
            start=1
        ):
            try:
                activity, grades, found_names = (
                    parse_activity(upload)
                )

                if not grades:
                    errors.append(
                        f"{upload.name}: "
                        "no se encontraron intentos válidos"
                    )

                elif normalize(activity) in seen:
                    errors.append(
                        f"{upload.name}: "
                        f"actividad duplicada ({activity})"
                    )

                else:
                    seen.add(normalize(activity))

                    activities.append(
                        (activity, grades)
                    )

                    names.update({
                        key: value
                        for key, value in found_names.items()
                        if key in grades and key not in names
                    })

            except Exception as exc:
                errors.append(
                    f"{upload.name}: {exc}"
                )

            progress.progress(
                i / len(uploads),
                text=(
                    f"Archivos revisados: "
                    f"{i} de {len(uploads)}"
                )
            )

        progress.empty()

        if activities:
            report, group_count = make_report(
                activities,
                names,
                title
            )

            st.session_state.report_bytes = report

            st.session_state.report_name = (
                clean_filename(title)
            )

            st.session_state.report_info = (
                len(activities),
                group_count
            )

        if errors:
            with st.expander(
                f"Avisos de archivos ({len(errors)})",
                expanded=True
            ):
                for message in errors:
                    st.warning(message)

        if not activities:
            st.error(
                "No se pudo generar el reporte: "
                "no hay actividades válidas."
            )

if st.session_state.report_bytes is not None:
    count, groups = st.session_state.report_info

    st.success(
        f"Reporte listo: {count} actividades "
        f"y {groups} grupos. "
        "No se muestran calificaciones en esta página."
    )

    st.download_button(
        "⬇️ Descargar reporte Excel",
        data=st.session_state.report_bytes,
        file_name=st.session_state.report_name,
        mime=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        use_container_width=True
    )

    if st.button(
        "Preparar otro reporte",
        use_container_width=True
    ):
        st.session_state.report_bytes = None
        st.session_state.report_name = None
        st.session_state.report_info = None

        st.rerun()
